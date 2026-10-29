"""Turn the 전국산업단지 현황통계 workbook (한국산업단지공단, quarterly) into scripts/static/industrial.csv.

Run on a PC with the downloaded file (needs openpyxl):
    python scripts/industrial_csv.py "전국산업단지현황통계_20260331.xlsx" 2026-03-31
One row per 상위단지 and 시군 (sheet '전국산업단지현황(상위단지)'); 노후 marks the 상위단지 in the 노후산업단지 sheet
(착공 후 20년 이상). Areas in 천㎡.
"""
import csv
import os
import sys

import openpyxl


def main(xlsx, as_of):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    old = {(str(r[1]).strip(), str(r[3]).strip()) for r in list(wb["부록4)노후산업단지"].iter_rows(values_only=True))[8:] if r[0] and r[1] and r[3]}
    rows = [r for r in list(wb["전국산업단지현황(상위단지)"].iter_rows(values_only=True))[6:] if r[0] and r[1] and r[3]]
    path = os.path.join(os.path.dirname(__file__), "static", "industrial.csv")
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["기준일", "유형", "시도", "시군구", "단지명", "조성상태", "관리면적_천m2", "입주업체", "가동업체", "노후"])
        for r in rows:
            name = " ".join(str(r[3]).split())
            w.writerow([as_of, r[0], str(r[1]).strip(), " ".join(str(r[2]).split()), name, r[4], r[6] or 0, r[12] or 0, r[13] or 0,
                        int((str(r[1]).strip(), str(r[3]).strip()) in old)])
    print(f"{path}: {len(rows)} rows, {sum((str(r[1]).strip(), str(r[3]).strip()) in old for r in rows)} 노후")


if __name__ == "__main__":
    main(*sys.argv[1:3])
